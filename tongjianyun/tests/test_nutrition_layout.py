import unittest
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from tongjianyun.ai_nutrition import NUTRIENTS, source_data
from tongjianyun.nutrition_layout import make_sheet, render_sheet
from tongjianyun.weekly_nutrition_sheet import workbook_bytes


def fixture(count=1):
    dates = [f"2026-06-{d}" for d in range(22, 27)]
    lines = [dict(meal_date=dates[0], meal_type="早餐", ingredient_code=str(i),
                  ingredient_name=f"大米{i}", uom="Gram", per_child_qty=25) for i in range(count)]
    data = source_data(lines, dates, {"week_start": dates[0], "plan_batch": "test", "week_name": None})
    result = {"categories": [{"food": i, "category": "细粮"} for i in range(count)],
              "nutrients": [{"key": key, "daily": [[10, 20]] * 5, "average": [10, 20], "basis": "合成测试数据"} for key in NUTRIENTS],
              "assumptions": ["合成数据"], "missing": []}
    return data, result


class LayoutTests(unittest.TestCase):
    def test_food_overflow_preserves_every_item_and_merged_cell_coverage(self):
        data, result = fixture(160)
        sheet = make_sheet(data, result)
        self.assertGreater(sheet["rows"], 36)
        names = [c["value"] for c in sheet["cells"] if c["kind"] == "food"]
        self.assertCountEqual(names, [f["name"] for f in data["foods"]])
        covered = set()
        for cell in sheet["cells"]:
            for r in range(cell["row"], cell["row"] + cell["rowspan"]):
                for c in range(cell["col"], cell["col"] + cell["colspan"]):
                    self.assertNotIn((r, c), covered)
                    covered.add((r, c))
        self.assertEqual(len(covered), sheet["rows"] * 21)

    def test_download_matches_page_cells_and_unknown_is_not_zero(self):
        data, result = fixture()
        result["nutrients"][0].update(daily=[None] * 5, average=None)
        cells = make_sheet(data, result)["cells"]
        ws = load_workbook(BytesIO(workbook_bytes(data, result))).active
        for item in cells:
            self.assertEqual(ws.cell(item["row"] + 2, item["col"] + 2).value or "", str(item["value"]))
        self.assertEqual(ws["T4"].value, "未提供")
        self.assertEqual(ws["R4"].value, "未设定")
        self.assertIn("B4:B11", [str(m) for m in ws.merged_cells.ranges])
        self.assertEqual(ws.page_setup.orientation, "landscape")

    def test_mixed_units_are_not_added_and_html_excel_injection_are_escaped(self):
        data, result = fixture(2)
        data["foods"][0].update(name="<script>alert(1)</script>", unit="Gram", daily="2")
        data["foods"][1].update(name="=1+1", unit="Millilitre", daily="3")
        html = render_sheet(data, result)
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        ws = load_workbook(BytesIO(workbook_bytes(data, result))).active
        values = [cell for row in ws for cell in row if cell.value is not None]
        self.assertTrue(any("2 g\n3 ml" in str(c.value) for c in values))
        self.assertEqual(next(c for c in values if c.value == "=1+1").data_type, "s")
        self.assertFalse(any(c.data_type == "f" for c in values))

    def test_source_only_page_has_no_nutritional_values(self):
        data, _ = fixture()
        cells = make_sheet(data)["cells"]
        self.assertEqual(next(c["value"] for c in cells if (c["row"], c["col"]) == (2, 18)), "待分析")
        self.assertIn("3—6岁混龄", render_sheet(data))
