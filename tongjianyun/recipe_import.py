"""I-ONE Agent assisted import for weekly Tongjianyun recipe workbooks."""

from __future__ import annotations

import json
import re
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any, Callable

import frappe
from frappe import _
from openpyxl import load_workbook

from ione_agent.structured import StructuredTaskClient


RECIPE_DOCTYPE = "Tongjianyun Recipe"
IMPORT_CACHE_PREFIX = "tongjianyun:recipe-import"
IMPORT_TTL_SECONDS = 3600
MAX_FILE_BYTES = 10 * 1024 * 1024
MEAL_SLOTS = {
    "早餐": "breakfast",
    "早点": "morningSnack",
    "午餐": "lunch",
    "午点": "snack",
    "晚餐": "dinner",
}
MEAL_LABELS = {slot: label for label, slot in MEAL_SLOTS.items()}
DAY_LABELS = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]
ROW_TYPES = {"主食品", "主食", "配餐", "带量", "食材", "用量"}
MASS_UNITS = {"kg": "kg", "千克": "kg", "公斤": "kg", "g": "g", "克": "g", "mg": "mg", "毫克": "mg"}
UNIT_ALIASES = {
    **MASS_UNITS,
    "ml": "ml",
    "毫升": "ml",
    "l": "L",
    "升": "L",
    "个": "个",
    "只": "只",
    "枚": "枚",
    "片": "片",
    "根": "根",
    "袋": "袋",
    "盒": "盒",
    "杯": "杯",
}
TOKEN_RE = re.compile(
    r"(?P<name>[^\d\s,，、;；:：]+(?:\s+[^\d\s,，、;；:：]+)*)?\s*"
    r"(?P<amount>\d+(?:\.\d+)?)\s*"
    r"(?P<unit>千克|公斤|毫克|毫升|kg|mg|ml|g|克|升|l|个|只|枚|片|根|袋|盒|杯)",
    re.IGNORECASE,
)


class RecipeImportError(ValueError):
    pass


def _require_recipe_write() -> None:
    if frappe.session.user == "Guest":
        frappe.throw(_("Not logged in."), frappe.AuthenticationError)
    if not frappe.has_permission(RECIPE_DOCTYPE, ptype="write"):
        frappe.throw(_("Not permitted."), frappe.PermissionError)


def _cache_key(import_id: str) -> str:
    return f"{IMPORT_CACHE_PREFIX}:{import_id}"


def _set_status(import_id: str, user: str, **updates: Any) -> dict[str, Any]:
    current = frappe.cache.get_value(
        _cache_key(import_id), user=user, expires=True, use_local_cache=False
    ) or {}
    current.update(updates)
    current["import_id"] = import_id
    current["updated_at"] = frappe.utils.now_datetime().isoformat()
    frappe.cache.set_value(
        _cache_key(import_id), current, user=user, expires_in_sec=IMPORT_TTL_SECONDS
    )
    return current


@frappe.whitelist()
def start_recipe_import(file_url: str) -> dict[str, Any]:
    _require_recipe_write()
    file_url = str(file_url or "").strip()
    if not file_url.lower().endswith(".xlsx"):
        frappe.throw("目前仅支持 .xlsx 食谱文件。")

    file_doc = frappe.get_doc("File", {"file_url": file_url})
    if not frappe.has_permission("File", ptype="read", doc=file_doc):
        frappe.throw(_("Not permitted."), frappe.PermissionError)
    if int(file_doc.file_size or 0) > MAX_FILE_BYTES:
        frappe.throw("食谱文件不能超过 10 MB。")

    user = frappe.session.user
    import_id = uuid.uuid4().hex
    _set_status(
        import_id,
        user,
        status="queued",
        progress=2,
        message="文件已上传，等待 I-ONE Agent 识别。",
        source_file=file_doc.file_name,
    )
    frappe.enqueue(
        "tongjianyun.recipe_import.run_recipe_import_job",
        queue="long",
        timeout=360,
        job_id=f"recipe-import-{import_id}",
        import_id=import_id,
        file_url=file_url,
        source_file=file_doc.file_name,
        user=user,
    )
    return {"import_id": import_id}


@frappe.whitelist()
def get_recipe_import_status(import_id: str) -> dict[str, Any]:
    if frappe.session.user == "Guest":
        frappe.throw(_("Not logged in."), frappe.AuthenticationError)
    import_id = str(import_id or "").strip()
    if not re.fullmatch(r"[0-9a-f]{32}", import_id):
        frappe.throw("导入任务编号无效。")
    result = frappe.cache.get_value(
        _cache_key(import_id),
        user=frappe.session.user,
        expires=True,
        use_local_cache=False,
    )
    if not result:
        frappe.throw("导入任务不存在或已过期。")
    return result


def run_recipe_import_job(import_id: str, file_url: str, source_file: str, user: str) -> None:
    frappe.set_user(user)
    try:
        _set_status(
            import_id,
            user,
            status="running",
            progress=8,
            message="正在读取 Excel 版式和餐次。",
        )
        file_doc = frappe.get_doc("File", {"file_url": file_url})
        content = file_doc.get_content()
        if isinstance(content, str):
            content = content.encode()
        if len(content) > MAX_FILE_BYTES:
            raise RecipeImportError("食谱文件不能超过 10 MB。")

        def report(progress: int, message: str) -> None:
            _set_status(
                import_id,
                user,
                status="running",
                progress=progress,
                message=message,
            )

        payload, warnings, summary = parse_recipe_workbook(
            content,
            source_file=source_file,
            import_id=import_id,
            progress_callback=report,
        )
        _set_status(
            import_id,
            user,
            status="completed",
            progress=100,
            message="I-ONE Agent 识别完成，请校对后保存草稿。",
            result={"payload": payload, "warnings": warnings, "summary": summary},
        )
    except Exception as exc:
        frappe.log_error(
            title=f"Recipe import failed: {source_file}",
            message=frappe.get_traceback(),
        )
        safe_message = str(exc) if isinstance(exc, RecipeImportError) else "识别失败，请检查文件格式后重试。"
        _set_status(
            import_id,
            user,
            status="failed",
            progress=100,
            message=safe_message,
        )


def parse_recipe_workbook(
    content: bytes,
    *,
    source_file: str,
    import_id: str,
    progress_callback: Callable[[int, str], None] | None = None,
    agent_client: StructuredTaskClient | None = None,
) -> tuple[dict[str, Any], list[str], dict[str, Any]]:
    """Extract a weekly recipe and use IONE Agent for semantic dish assignment."""

    report = progress_callback or (lambda _progress, _message: None)
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        raise RecipeImportError("无法读取 Excel，请确认文件未损坏且格式为 .xlsx。") from exc

    sheet = _select_sheet(workbook)
    title = _find_title(sheet)
    year = _find_year(title, sheet)
    date_columns = _find_date_columns(sheet, year)
    if not date_columns:
        raise RecipeImportError("未识别到日期列，请保留“月/日/星期”表头。")
    if len(date_columns) > 7:
        raise RecipeImportError("日期列超过 7 天，请一次导入一周食谱。")

    warnings: list[str] = []
    extracted_days = _extract_days(sheet, date_columns, warnings)
    if not any(day["portions"] for day in extracted_days):
        raise RecipeImportError("未识别到早餐、早点、午餐、午点或晚餐内容。")
    report(18, f"已识别 {len(extracted_days)} 个日期，正在调用 I-ONE Agent。")

    client = agent_client or StructuredTaskClient.from_frappe_config()
    resolved: dict[int, list[dict[str, Any]]] = {}
    worker_count = min(5, max(1, len(extracted_days)))
    with ThreadPoolExecutor(max_workers=worker_count) as pool:
        futures = {
            pool.submit(_resolve_day_with_agent, client, day): index
            for index, day in enumerate(extracted_days)
        }
        completed = 0
        for future in as_completed(futures):
            index = futures[future]
            day = extracted_days[index]
            try:
                portions, day_warnings = future.result()
            except Exception:
                portions, day_warnings = _fallback_day(day, "Agent 返回异常，已按原文合并导入")
            resolved[index] = portions
            warnings.extend(day_warnings)
            completed += 1
            progress = 20 + round(65 * completed / len(extracted_days))
            report(progress, f"I-ONE Agent 已完成 {completed}/{len(extracted_days)} 天的语义识别。")

    days: list[dict[str, Any]] = []
    for index, extracted in enumerate(extracted_days):
        days.append(
            {
                "id": f"DAY-{index + 1}",
                "date": extracted["date"].isoformat(),
                "day": extracted["day_label"],
                "portions": resolved[index],
                "version": 1,
            }
        )

    week_start = min(day["date"] for day in extracted_days)
    week_end = max(day["date"] for day in extracted_days)
    recipe_id = f"RECIPE-AI-{week_start:%Y%m%d}-{import_id[:8].upper()}"
    payload = {
        "recipe": {
            "recipeId": recipe_id,
            "title": title or f"{week_start:%Y年%m月%d日}周食谱",
            "weekStart": week_start.isoformat(),
            "weekEnd": week_end.isoformat(),
            "sourceFileName": PurePosixPath(source_file or "食谱.xlsx").name,
            "parser": "I-ONE Agent Structured Import",
            "relationSource": f"ione-agent:{import_id}",
            "workflowStatus": "草稿",
            "allStudentGroups": True,
            "studentGroups": [],
        },
        "days": days,
    }
    dish_count = sum(len(portion["dishes"]) for day in days for portion in day["portions"])
    ingredient_count = sum(
        len(portion["dishIngredientRows"]) for day in days for portion in day["portions"]
    )
    summary = {
        "sheet": sheet.title,
        "day_count": len(days),
        "meal_count": sum(len(day["portions"]) for day in days),
        "dish_count": dish_count,
        "ingredient_count": ingredient_count,
        "warning_count": len(warnings),
    }
    report(92, "正在校验菜品、食材和用量关联。")
    return payload, warnings, summary


def _select_sheet(workbook):
    for sheet in workbook.worksheets:
        if any(cell.value not in (None, "") for row in sheet.iter_rows() for cell in row):
            return sheet
    raise RecipeImportError("Excel 中没有可导入的数据。")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def _find_title(sheet) -> str:
    candidates: list[str] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8)):
        for cell in row:
            text = _clean(cell.value)
            if text:
                candidates.append(text)
                if "食谱" in text and ("年" in text or "周" in text):
                    return text
    return max(candidates, key=len, default="")


def _find_year(title: str, sheet) -> int:
    match = re.search(r"(20\d{2})\s*年", title)
    if match:
        return int(match.group(1))
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8)):
        for cell in row:
            if isinstance(cell.value, (datetime, date)):
                return cell.value.year
    return frappe.utils.now_datetime().year


def _parse_header_date(value: Any, year: int) -> tuple[date, str] | None:
    if isinstance(value, datetime):
        return value.date(), DAY_LABELS[value.weekday()]
    if isinstance(value, date):
        return value, DAY_LABELS[value.weekday()]
    text = _clean(value)
    match = re.search(r"(?:(20\d{2})\s*年)?\s*(\d{1,2})\s*月\s*(\d{1,2})\s*[日号]?", text)
    if not match:
        return None
    parsed = date(int(match.group(1) or year), int(match.group(2)), int(match.group(3)))
    weekday = re.search(r"(?:周|星期)\s*([一二三四五六日天])", text)
    label = f"星期{weekday.group(1).replace('天', '日')}" if weekday else DAY_LABELS[parsed.weekday()]
    return parsed, label


def _find_date_columns(sheet, year: int) -> list[dict[str, Any]]:
    candidates: list[dict[int, dict[str, Any]]] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8)):
        found: dict[int, dict[str, Any]] = {}
        for cell in row:
            parsed = _parse_header_date(cell.value, year)
            if parsed:
                found[cell.column] = {"column": cell.column, "date": parsed[0], "day_label": parsed[1]}
        if found:
            candidates.append(found)
    if not candidates:
        return []
    # A title often contains the first date too.  The actual header is the row
    # containing the greatest number of distinct daily columns.
    best = max(candidates, key=lambda values: len(values))
    return sorted(best.values(), key=lambda item: item["date"])


def _extract_days(sheet, date_columns: list[dict[str, Any]], warnings: list[str]) -> list[dict[str, Any]]:
    result = [{**column, "portions": {}} for column in date_columns]
    current_slot = ""
    for row_index in range(1, sheet.max_row + 1):
        meal_text = _clean(sheet.cell(row_index, 1).value).replace(" ", "")
        if meal_text in MEAL_SLOTS:
            current_slot = MEAL_SLOTS[meal_text]
        row_type = _clean(sheet.cell(row_index, 2).value).replace(" ", "")
        if not current_slot or row_type not in ROW_TYPES:
            continue

        for day in result:
            value = _clean(sheet.cell(row_index, day["column"]).value)
            if not value:
                continue
            portion = day["portions"].setdefault(
                current_slot,
                {"slot": current_slot, "label": MEAL_LABELS[current_slot], "dish_parts": [], "quantity_parts": []},
            )
            if row_type in {"带量", "食材", "用量"}:
                portion["quantity_parts"].append(value)
            else:
                portion["dish_parts"].append(value)

    extracted: list[dict[str, Any]] = []
    for day in result:
        portions: list[dict[str, Any]] = []
        for slot in MEAL_LABELS:
            source = day["portions"].get(slot)
            if not source:
                continue
            dish_text = " ".join(source["dish_parts"]).strip()
            quantity_text = " ".join(source["quantity_parts"]).strip()
            ingredients = _parse_ingredients(quantity_text)
            if quantity_text and not ingredients:
                warnings.append(f"{day['date']} {MEAL_LABELS[slot]}：带量文本未识别为数值，已保留原文。")
            portions.append(
                {
                    "slot": slot,
                    "label": MEAL_LABELS[slot],
                    "dish_text": dish_text,
                    "quantity_text": quantity_text,
                    "ingredients": ingredients,
                }
            )
        extracted.append({**day, "portions": portions})
    return extracted


def _parse_ingredients(text: str) -> list[dict[str, Any]]:
    ingredients: list[dict[str, Any]] = []
    for index, match in enumerate(TOKEN_RE.finditer(text)):
        unit_source = match.group("unit")
        unit = UNIT_ALIASES.get(unit_source.lower(), UNIT_ALIASES.get(unit_source, unit_source))
        name = _clean(match.group("name"))
        ingredients.append(
            {
                "index": index,
                "name": name,
                "amount": float(match.group("amount")),
                "unit": unit,
                "raw": match.group(0).strip(),
            }
        )
    return ingredients


AGENT_SYSTEM_PROMPT = """你是 I-ONE Agent 的幼儿园周食谱结构化识别器。
输入是一整天各餐次的原始菜名文本和已由程序提取的带量食材。
请只返回 JSON 对象：{"portions":[{"slot":"breakfast","dishes":["菜名"],"assignments":[0],"inferred":false}]}。
规则：
1. portions 必须逐项覆盖输入中的餐次，slot 原样返回。
2. 将粘连或用空格分隔的菜名拆成自然菜品；不得改写、补充或遗漏原文字符（空格和标点除外）。
3. assignments 与 ingredients 等长，每个整数表示该食材属于 dishes 的哪个下标。必须逐个按食材名称的烹饪语义归属，不能按位置轮流分配。
   示例：dishes=["纯牛奶","蓝莓切片面包","水煮鹌鹑蛋"]，ingredients=["纯牛奶","切片面包","蓝莓果酱","鹌鹑蛋"]，assignments 必须是 [0,1,1,2]。
4. 菜名为空但存在食材时，可根据食材推断一个简短菜名，并将 inferred 设为 true；不要推断用量。
5. 菜名和食材均为空时 dishes 返回空数组。不要输出解释或 Markdown。"""


def _resolve_day_with_agent(
    client: StructuredTaskClient, day: dict[str, Any]
) -> tuple[list[dict[str, Any]], list[str]]:
    agent_input = {
        "date": day["date"].isoformat(),
        "portions": [
            {
                "slot": portion["slot"],
                "label": portion["label"],
                "dish_text": portion["dish_text"],
                "ingredients": portion["ingredients"],
            }
            for portion in day["portions"]
        ],
    }
    result = client.complete_json(system_prompt=AGENT_SYSTEM_PROMPT, user_payload=agent_input)
    result_by_slot = {
        str(item.get("slot") or ""): item
        for item in result.get("portions", [])
        if isinstance(item, dict)
    }

    warnings: list[str] = []
    staged: list[dict[str, Any]] = []
    for source in day["portions"]:
        agent_portion = result_by_slot.get(source["slot"], {})
        dishes = [_clean(value) for value in agent_portion.get("dishes", []) if _clean(value)]
        inferred = bool(agent_portion.get("inferred"))
        if source["dish_text"]:
            if not dishes or _canonical("".join(dishes)) != _canonical(source["dish_text"]):
                dishes = [source["dish_text"]]
                warnings.append(
                    f"{day['date']} {source['label']}：Agent 拆分未通过原文校验，已按原文合并。"
                )
                inferred = False
        elif source["ingredients"]:
            if not dishes:
                dishes = [f"未命名菜品（{source['label']}）"]
            inferred = True
        else:
            dishes = []

        if inferred:
            warnings.append(f"{day['date']} {source['label']}：原表菜名为空，Agent 已根据食材推断，请重点校对。")

        initial_assignments = _normalize_assignments(
            agent_portion.get("assignments", []),
            ingredient_count=len(source["ingredients"]),
            dish_count=len(dishes),
        )
        staged.append(
            {
                "source": source,
                "dishes": dishes,
                "initial_assignments": initial_assignments,
            }
        )

    repair_candidates = [
        item
        for item in staged
        if len(item["dishes"]) > 1 and item["source"]["ingredients"]
    ]
    repaired: dict[str, list[int]] = {}
    if repair_candidates:
        try:
            repaired = _repair_ingredient_links(client, repair_candidates)
        except Exception:
            warnings.append(f"{day['date']}：食材归属复核未完成，已使用首轮识别结果。")

    normalized: list[dict[str, Any]] = []
    for item in staged:
        source = item["source"]
        dishes = item["dishes"]
        if len(dishes) == 1:
            assignments = [0] * len(source["ingredients"])
        else:
            assignments = repaired.get(source["slot"]) or item["initial_assignments"]
        if assignments is None:
            assignments = _heuristic_assignments(source["ingredients"], dishes)
            if source["ingredients"]:
                warnings.append(f"{day['date']} {source['label']}：部分食材归属需人工校对。")
        normalized.append(_build_portion(source, dishes, assignments))
    return normalized, warnings


REPAIR_SYSTEM_PROMPT = """你是 I-ONE Agent 食材归属校验器。只返回 JSON：
{"portions":[{"slot":"lunch","links":["菜名","菜名"]}]}。
每个 links 必须与输入 ingredients 等长、顺序一致；每个值必须原样选自同餐次 dishes。
按实际烹饪语义逐个判断。例：dishes=[纯牛奶,蓝莓切片面包,水煮鹌鹑蛋]，ingredients=[纯牛奶,切片面包,蓝莓果酱,鹌鹑蛋]，links=[纯牛奶,蓝莓切片面包,蓝莓切片面包,水煮鹌鹑蛋]。
不要返回序号，不要解释。"""


def _repair_ingredient_links(
    client: StructuredTaskClient, candidates: list[dict[str, Any]]
) -> dict[str, list[int]]:
    payload = {
        "portions": [
            {
                "slot": item["source"]["slot"],
                "dishes": item["dishes"],
                "ingredients": [ingredient["name"] for ingredient in item["source"]["ingredients"]],
            }
            for item in candidates
        ]
    }
    result = client.complete_json(system_prompt=REPAIR_SYSTEM_PROMPT, user_payload=payload)
    result_by_slot = {
        str(item.get("slot") or ""): item
        for item in result.get("portions", [])
        if isinstance(item, dict)
    }
    repaired: dict[str, list[int]] = {}
    for candidate in candidates:
        source = candidate["source"]
        dishes = candidate["dishes"]
        links = result_by_slot.get(source["slot"], {}).get("links", [])
        if not isinstance(links, list) or len(links) != len(source["ingredients"]):
            continue
        if any(link not in dishes for link in links):
            continue
        repaired[source["slot"]] = [dishes.index(link) for link in links]
    return repaired


def _normalize_assignments(
    values: Any, *, ingredient_count: int, dish_count: int
) -> list[int] | None:
    if ingredient_count == 0:
        return []
    if dish_count == 1:
        return [0] * ingredient_count
    if not isinstance(values, list) or len(values) != ingredient_count or dish_count < 1:
        return None
    assignments: list[int] = []
    for value in values:
        try:
            index = int(value)
        except (TypeError, ValueError):
            return None
        if index < 0 or index >= dish_count:
            return None
        assignments.append(index)
    return assignments


def _heuristic_assignments(
    ingredients: list[dict[str, Any]], dishes: list[str]
) -> list[int]:
    if not dishes:
        return [0] * len(ingredients)
    canonical_dishes = [_canonical(dish) for dish in dishes]
    assignments: list[int] = []
    for ingredient in ingredients:
        name = _canonical(ingredient.get("name") or "")
        matches = [
            index
            for index, dish in enumerate(canonical_dishes)
            if name and (name in dish or dish in name)
        ]
        assignments.append(matches[0] if len(matches) == 1 else 0)
    return assignments


def _canonical(value: str) -> str:
    return re.sub(r"[\s,，、;；:：/|·。．（）()【】\[\]—_-]+", "", value or "")


def _build_portion(
    source: dict[str, Any], dishes: list[str], assignments: list[int]
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for ingredient, dish_index in zip(source["ingredients"], assignments):
        dish_name = dishes[dish_index] if dishes else f"未命名菜品（{source['label']}）"
        ingredient_name = ingredient["name"] or (dish_name if len(dishes) == 1 else dish_name)
        amount = ingredient["amount"]
        unit = ingredient["unit"]
        row = {
            "dishName": dish_name,
            "ingredient": ingredient_name,
            "amount": amount,
            "unit": unit,
        }
        if unit == "g":
            row["gramsPerChild"] = amount
        elif unit == "kg":
            row["gramsPerChild"] = amount * 1000
        elif unit == "mg":
            row["gramsPerChild"] = amount / 1000
        rows.append(row)
    return {
        "slot": source["slot"],
        "label": source["label"],
        "dishes": dishes,
        "amountPerChild": source["quantity_text"],
        "totalAmount": "",
        "dishIngredientRows": rows,
    }


def _fallback_day(day: dict[str, Any], reason: str) -> tuple[list[dict[str, Any]], list[str]]:
    portions: list[dict[str, Any]] = []
    warnings: list[str] = []
    for source in day["portions"]:
        dishes = [source["dish_text"]] if source["dish_text"] else []
        if not dishes and source["ingredients"]:
            dishes = [f"未命名菜品（{source['label']}）"]
        assignments = [0] * len(source["ingredients"])
        portions.append(_build_portion(source, dishes, assignments))
        warnings.append(f"{day['date']} {source['label']}：{reason}，请校对。")
    return portions, warnings
